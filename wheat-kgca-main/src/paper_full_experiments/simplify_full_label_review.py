from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = PROJECT_ROOT / "outputs" / "paper_full_experiments" / "full_label_review" / "full_label_review_all.csv"
DEFAULT_OUT_DIR = PROJECT_ROOT / "outputs" / "paper_full_experiments" / "full_label_review_simple"

LABELS = [
    "seedling",
    "tillering",
    "stem_elongation",
    "flag_leaf_booting",
    "heading",
    "flowering",
    "grain_filling",
    "maturity",
]

LABEL_CN = {
    "seedling": "苗期",
    "tillering": "分蘖期",
    "stem_elongation": "拔节期",
    "flag_leaf_booting": "旗叶孕穗期",
    "heading": "抽穗期",
    "flowering": "开花期",
    "grain_filling": "灌浆期",
    "maturity": "成熟期",
}


def build_simple_table(source: Path) -> pd.DataFrame:
    df = pd.read_csv(source, low_memory=False)
    out = pd.DataFrame()
    out["review_action"] = ""
    out["new_label"] = ""
    out["note"] = ""
    out["image_id"] = df["image_id"]
    out["original_label"] = df["stage_id"]
    out["original_label_cn"] = df["stage_cn"]
    out["image_path"] = df["image_path"]
    out["source_path"] = df.get("source_path", "")
    out["original_filename"] = df.get("original_filename", "")
    out["review_priority"] = df.get("review_priority", "")
    out["review_reason"] = df.get("review_reason", "")
    out["source_type"] = df.get("source_type", "")
    out["convnext_pred"] = df.get("convnext_pred", "")
    out["convnext_confidence"] = df.get("convnext_confidence", "")
    out["convnext_correct"] = df.get("convnext_correct", "")
    out["vlm_predicted_stage"] = df.get("vlm_predicted_stage", "")
    out["vlm_second_candidate"] = df.get("vlm_second_candidate", "")
    out["vlm_confidence"] = df.get("vlm_confidence", "")
    out["vlm_plant_color"] = df.get("vlm_plant_color", "")
    out["vlm_spike_status"] = df.get("vlm_spike_status", "")
    out["vlm_anther_visible"] = df.get("vlm_anther_visible", "")
    out["vlm_grain_fullness"] = df.get("vlm_grain_fullness", "")
    out["vlm_yellowing_degree"] = df.get("vlm_yellowing_degree", "")
    out["vlm_short_reason"] = df.get("vlm_short_reason", "")
    out["vlm_uncertainty_reason"] = df.get("vlm_uncertainty_reason", "")
    out["boundary_hint"] = (
        df.get("original_boundary_pair_hint", "").fillna("").astype(str)
        + ";"
        + df.get("convnext_boundary_pair_hint", "").fillna("").astype(str)
        + ";"
        + df.get("vlm_boundary_pair_hint", "").fillna("").astype(str)
    ).str.strip(";")
    out["old_split_trace_only"] = df.get("split", "")
    out["target_path"] = df.get("target_path", "")

    out = out.sort_values(["review_priority", "original_label", "image_id"], kind="stable")
    return out


def write_xlsx(df: pd.DataFrame, out_path: Path) -> bool:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return False

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        guide = pd.DataFrame(
            [
                {
                    "column": "review_action",
                    "how_to_fill": "Fill keep if the original label is correct; fill relabel if it is wrong.",
                },
                {
                    "column": "new_label",
                    "how_to_fill": "Only fill this when review_action is relabel. Use one of the eight English labels.",
                },
                {
                    "column": "note",
                    "how_to_fill": "Optional short reason, especially for maturity/grain_filling and other boundary cases.",
                },
                {
                    "column": "old_split_trace_only",
                    "how_to_fill": "Traceability only. Do not use it as review criterion. Final data will be re-split after relabeling.",
                },
            ]
        )
        labels = pd.DataFrame({"label": LABELS, "label_cn": [LABEL_CN[x] for x in LABELS]})
        guide.to_excel(writer, sheet_name="guide", index=False)
        labels.to_excel(writer, sheet_name="label_options", index=False)
        df.to_excel(writer, sheet_name="all_review", index=False)

    wb = load_workbook(out_path)
    ws = wb["all_review"]
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for col in ["A", "B", "C"]:
        ws[f"{col}1"].fill = editable_fill
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    action_validation = DataValidation(type="list", formula1='"keep,relabel"', allow_blank=True)
    label_validation = DataValidation(type="list", formula1=f'"{",".join(LABELS)}"', allow_blank=True)
    ws.add_data_validation(action_validation)
    ws.add_data_validation(label_validation)
    last_row = ws.max_row
    action_validation.add(f"A2:A{last_row}")
    label_validation.add(f"B2:B{last_row}")

    widths = {
        "A": 14,
        "B": 20,
        "C": 30,
        "D": 28,
        "E": 22,
        "F": 16,
        "G": 80,
        "H": 80,
        "I": 30,
        "J": 14,
        "K": 40,
        "R": 22,
        "S": 20,
        "X": 60,
        "Y": 60,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(out_path)
    return True


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create a simplified full-label review workbook.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    df = build_simple_table(args.source)

    csv_path = args.out_dir / "full_label_review_simple_all.csv"
    xlsx_path = args.out_dir / "full_label_review_simple_all.xlsx"
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    wrote_xlsx = write_xlsx(df, xlsx_path)

    print(f"Wrote CSV: {csv_path}")
    if wrote_xlsx:
        print(f"Wrote XLSX: {xlsx_path}")
    else:
        print("openpyxl is not installed; XLSX skipped.")
    print(f"Rows: {len(df)}")


if __name__ == "__main__":
    main()
